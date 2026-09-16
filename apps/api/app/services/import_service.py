import csv
import io
import datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional
import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.config import settings
from app.core.errors import EntityNotFoundException, FileValidationException
from app.db.models.account import Account
from app.db.models.category import Category
from app.db.models.transaction import Transaction, TransactionType
from app.repositories.account_repo import AccountRepository
from app.repositories.transaction_repo import TransactionRepository
from app.schemas.imports import (
    ColumnMapping,
    ImportPreviewResponse,
    ImportSummaryResponse,
    ImportErrorDetail,
)


class ImportService:
    def __init__(self, session: AsyncSession):
        self.session = session
        self.account_repo = AccountRepository(session)
        self.tx_repo = TransactionRepository(session)

    def _parse_tabular_file(self, content: bytes, filename: str) -> List[Dict[str, Any]]:
        # Validate file size
        if len(content) > settings.MAX_UPLOAD_SIZE_BYTES:
            raise FileValidationException(
                f"File size ({len(content)} bytes) exceeds the maximum allowed limit of {settings.MAX_UPLOAD_SIZE_BYTES} bytes."
            )

        rows = []
        if filename.lower().endswith(".csv"):
            text = content.decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(io.StringIO(text))
            for row in reader:
                rows.append({k.strip(): (v.strip() if v else "") for k, v in row.items() if k})
        elif filename.lower().endswith((".xlsx", ".xls")):
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            ws = wb.active
            iter_rows = list(ws.iter_rows(values_only=True))
            if not iter_rows:
                return []
            header = [str(col).strip() if col is not None else f"col_{i}" for i, col in enumerate(iter_rows[0])]
            for r in iter_rows[1:]:
                if any(r):
                    rows.append({header[i]: (str(val).strip() if val is not None else "") for i, val in enumerate(r) if i < len(header)})
        else:
            raise FileValidationException("Unsupported file format. Please upload CSV or Excel (.xlsx) file.")
        return rows

    def _detect_columns(self, headers: List[str]) -> ColumnMapping:
        mapping = ColumnMapping()
        for col in headers:
            lower = col.lower()
            if not mapping.date and any(w in lower for w in ["date", "time", "posted"]):
                mapping.date = col
            elif not mapping.debit and any(w in lower for w in ["debit", "withdrawal", "spent"]):
                mapping.debit = col
            elif not mapping.credit and any(w in lower for w in ["credit", "deposit", "received"]):
                mapping.credit = col
            elif not mapping.amount and any(w in lower for w in ["amount", "total", "net"]):
                mapping.amount = col
            elif not mapping.description and any(w in lower for w in ["description", "narration", "particulars", "memo", "title"]):
                mapping.description = col
            elif not mapping.merchant and any(w in lower for w in ["merchant", "payee", "vendor"]):
                mapping.merchant = col
            elif not mapping.category and any(w in lower for w in ["category", "type", "tag"]):
                mapping.category = col
        return mapping

    async def preview_file(
        self, content: bytes, filename: str, user_id: str, account_id: Optional[str] = None
    ) -> ImportPreviewResponse:
        rows = self._parse_tabular_file(content, filename)
        if not rows:
            raise FileValidationException("The uploaded file is empty.")

        headers = list(rows[0].keys())
        mapping = self._detect_columns(headers)

        return ImportPreviewResponse(
            filename=filename,
            total_rows=len(rows),
            detected_columns=headers,
            suggested_mapping=mapping,
            preview_rows=rows[:10],
        )

    async def execute_import(
        self,
        content: bytes,
        filename: str,
        user_id: str,
        account_id: str,
        mapping: ColumnMapping,
        skip_duplicates: bool = True,
    ) -> ImportSummaryResponse:
        account = await self.account_repo.get_by_id(account_id)
        if not account or account.user_id != user_id:
            raise EntityNotFoundException("Account", account_id)

        rows = self._parse_tabular_file(content, filename)
        imported = 0
        skipped = 0
        failed = 0
        errors: List[ImportErrorDetail] = []

        balance_delta = Decimal("0.00")

        for idx, row in enumerate(rows, start=1):
            try:
                # 1. Parse Date
                date_val = None
                date_str = row.get(mapping.date or "", "")
                for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"):
                    try:
                        date_val = datetime.datetime.strptime(date_str.split()[0], fmt).date()
                        break
                    except Exception:
                        continue
                if not date_val:
                    date_val = datetime.date.today()

                # 2. Parse Amount & Type with INR/USD currency symbol stripping
                amount = Decimal("0.00")
                tx_type = TransactionType.EXPENSE
                if mapping.amount and row.get(mapping.amount):
                    raw_amt = (
                        row[mapping.amount]
                        .replace("₹", "")
                        .replace("$", "")
                        .replace(",", "")
                        .strip()
                    )
                    amt_val = Decimal(raw_amt)
                    if amt_val < Decimal("0.00"):
                        amount = abs(amt_val)
                        tx_type = TransactionType.EXPENSE
                    else:
                        amount = amt_val
                        tx_type = TransactionType.INCOME if "income" in row.get(mapping.type or "", "").lower() else TransactionType.EXPENSE
                elif mapping.debit and row.get(mapping.debit):
                    raw_amt = row[mapping.debit].replace("₹", "").replace("$", "").replace(",", "").strip()
                    amount = Decimal(raw_amt)
                    tx_type = TransactionType.EXPENSE
                elif mapping.credit and row.get(mapping.credit):
                    raw_amt = row[mapping.credit].replace("₹", "").replace("$", "").replace(",", "").strip()
                    amount = Decimal(raw_amt)
                    tx_type = TransactionType.INCOME

                # 3. Description & Merchant
                desc = row.get(mapping.description or "", "Imported Transaction").strip() or "Imported Transaction"
                merchant = row.get(mapping.merchant or "", "").strip() or None

                # 4. Duplicate Hash Prevention
                import_hash = Transaction.generate_hash(
                    account_id=account_id,
                    tx_date=date_val,
                    amount=amount,
                    desc=desc
                )

                stmt = select(Transaction).where(
                    Transaction.user_id == user_id,
                    Transaction.import_hash == import_hash
                )
                existing = (await self.session.execute(stmt)).scalar_one_or_none()
                if existing:
                    if skip_duplicates:
                        skipped += 1
                        continue

                # 5. Create Transaction
                tx = Transaction(
                    user_id=user_id,
                    account_id=account_id,
                    amount=amount,
                    currency=account.currency or "INR",
                    transaction_type=tx_type,
                    transaction_date=date_val,
                    description=desc,
                    merchant_name=merchant,
                    import_hash=import_hash,
                    is_cleared=True
                )
                self.session.add(tx)
                imported += 1

                if tx_type == TransactionType.INCOME:
                    balance_delta += amount
                elif tx_type == TransactionType.EXPENSE:
                    balance_delta -= amount

            except Exception as e:
                failed += 1
                errors.append(ImportErrorDetail(row_number=idx, reason=str(e), data=row))

        # Update account balance in batch safely
        account.current_balance = Decimal(str(account.current_balance)) + balance_delta
        await self.session.commit()

        return ImportSummaryResponse(
            total_rows_processed=len(rows),
            imported_count=imported,
            skipped_duplicates_count=skipped,
            failed_count=failed,
            errors=errors[:10],
        )
