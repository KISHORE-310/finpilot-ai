import csv
import io
import os
import datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Set
import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.config import settings
from app.core.errors import EntityNotFoundException, FileValidationException
from app.db.models.account import Account
from app.db.models.category import Category
from app.db.models.transaction import Transaction, TransactionType
from app.repositories.account_repo import AccountRepository
from app.repositories.transaction_repo import TransactionRepository
from app.schemas.imports import (
    ColumnMapping,
    ImportPreviewResponse,
    ImportSummaryResponse,
    ImportErrorDetail,
)


class ImportService:
    ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".xls"}

    def __init__(self, session: AsyncSession):
        self.session = session
        self.account_repo = AccountRepository(session)
        self.tx_repo = TransactionRepository(session)

    def _sanitize_filename(self, filename: str) -> str:
        """Sanitize filename to prevent path traversal vulnerabilities."""
        clean = os.path.basename(filename.strip().replace("\\", "/"))
        return clean or "uploaded_file.csv"

    def _parse_tabular_file(self, content: bytes, filename: str) -> List[Dict[str, Any]]:
        # 1. Validate non-empty content
        if not content or len(content) == 0:
            raise FileValidationException("The uploaded file is empty (0 bytes).")

        # 2. Validate file size
        if len(content) > settings.MAX_UPLOAD_SIZE_BYTES:
            raise FileValidationException(
                f"File size ({len(content)} bytes) exceeds the maximum allowed limit of {settings.MAX_UPLOAD_SIZE_BYTES} bytes."
            )

        # 3. Validate file extension
        sanitized_filename = self._sanitize_filename(filename)
        _, ext = os.path.splitext(sanitized_filename.lower())
        if ext not in self.ALLOWED_EXTENSIONS:
            raise FileValidationException(
                f"Unsupported file format '{ext}'. Please upload a CSV (.csv) or Excel (.xlsx, .xls) file."
            )

        rows = []
        if ext == ".csv":
            # Attempt decoding with UTF-8 BOM, UTF-8, then fallback to Latin-1
            text = None
            for encoding in ("utf-8-sig", "utf-8", "latin-1"):
                try:
                    text = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue

            if text is None:
                raise FileValidationException("Could not decode CSV file. Please ensure valid text encoding (UTF-8).")

            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames:
                raise FileValidationException("CSV file contains no valid headers or data columns.")

            for row in reader:
                clean_row = {k.strip(): (v.strip() if v else "") for k, v in row.items() if k}
                if any(clean_row.values()):
                    rows.append(clean_row)

        elif ext in (".xlsx", ".xls"):
            try:
                wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
                ws = wb.active
                if ws is None:
                    raise FileValidationException("Excel workbook contains no active worksheet.")
                iter_rows = list(ws.iter_rows(values_only=True))
                if not iter_rows:
                    return []
                header = [str(col).strip() if col is not None else f"col_{i}" for i, col in enumerate(iter_rows[0])]
                for r in iter_rows[1:]:
                    if any(r):
                        row_dict = {header[i]: (str(val).strip() if val is not None else "") for i, val in enumerate(r) if i < len(header)}
                        if any(row_dict.values()):
                            rows.append(row_dict)
            except Exception as e:
                raise FileValidationException(f"Failed to parse Excel workbook: {str(e)}")

        return rows

    def _detect_columns(self, headers: List[str]) -> ColumnMapping:
        mapping = ColumnMapping()
        for col in headers:
            lower = col.lower()
            if not mapping.date and any(w in lower for w in ["date", "time", "posted", "txn_date", "transaction date"]):
                mapping.date = col
            elif not mapping.debit and any(w in lower for w in ["debit", "withdrawal", "spent", "dr", "withdraw"]):
                mapping.debit = col
            elif not mapping.credit and any(w in lower for w in ["credit", "deposit", "received", "cr"]):
                mapping.credit = col
            elif not mapping.amount and any(w in lower for w in ["amount", "total", "net", "txn_amount", "transaction amount"]):
                mapping.amount = col
            elif not mapping.description and any(w in lower for w in ["description", "narration", "particulars", "memo", "title", "details", "remark"]):
                mapping.description = col
            elif not mapping.merchant and any(w in lower for w in ["merchant", "payee", "vendor", "party"]):
                mapping.merchant = col
            elif not mapping.category and any(w in lower for w in ["category", "type", "tag"]):
                mapping.category = col
        return mapping

    async def preview_file(
        self, content: bytes, filename: str, user_id: str, account_id: Optional[str] = None
    ) -> ImportPreviewResponse:
        sanitized_filename = self._sanitize_filename(filename)

        if account_id:
            account = await self.account_repo.get_by_id(account_id)
            if not account or account.user_id != user_id:
                raise EntityNotFoundException("Account", account_id)

        rows = self._parse_tabular_file(content, sanitized_filename)
        if not rows:
            raise FileValidationException("The uploaded file contains no data rows.")

        headers = list(rows[0].keys())
        mapping = self._detect_columns(headers)

        return ImportPreviewResponse(
            filename=sanitized_filename,
            total_rows=len(rows),
            detected_columns=headers,
            suggested_mapping=mapping,
            preview_rows=rows[:10],
        )

    async def execute_import(
        self,
        content: bytes,
        filename: str,
        user_id: str,
        account_id: str,
        mapping: ColumnMapping,
        skip_duplicates: bool = True,
    ) -> ImportSummaryResponse:
        sanitized_filename = self._sanitize_filename(filename)

        # Validate account existence and user ownership
        account = await self.account_repo.get_by_id(account_id)
        if not account or account.user_id != user_id:
            raise EntityNotFoundException("Account", account_id)

        rows = self._parse_tabular_file(content, sanitized_filename)
        if not rows:
            raise FileValidationException("The uploaded file contains no data rows to import.")

        imported = 0
        skipped = 0
        failed = 0
        errors: List[ImportErrorDetail] = []
        seen_hashes_in_batch: Set[str] = set()
        balance_delta = Decimal("0.00")

        for idx, row in enumerate(rows, start=1):
            try:
                # 1. Parse Date
                date_val = None
                date_str = row.get(mapping.date or "", "").strip()
                if date_str:
                    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y", "%d.%m.%Y", "%Y.%m.%d"):
                        try:
                            date_val = datetime.datetime.strptime(date_str.split()[0], fmt).date()
                            break
                        except (ValueError, TypeError):
                            continue
                if not date_val:
                    date_val = datetime.date.today()

                # 2. Parse Amount & Type with INR/USD currency symbol stripping
                amount = Decimal("0.00")
                tx_type = TransactionType.EXPENSE

                if mapping.amount and row.get(mapping.amount):
                    raw_amt = (
                        row[mapping.amount]
                        .replace("₹", "")
                        .replace("$", "")
                        .replace("€", "")
                        .replace("£", "")
                        .replace(",", "")
                        .strip()
                    )
                    try:
                        amt_val = Decimal(raw_amt)
                    except (InvalidOperation, ValueError):
                        failed += 1
                        errors.append(ImportErrorDetail(row_number=idx, reason=f"Invalid numeric amount '{raw_amt}'", data=row))
                        continue

                    if amt_val < Decimal("0.00"):
                        amount = abs(amt_val)
                        tx_type = TransactionType.EXPENSE
                    else:
                        amount = amt_val
                        raw_type_str = row.get(mapping.type or "", "").lower()
                        tx_type = TransactionType.INCOME if "income" in raw_type_str or "credit" in raw_type_str or "deposit" in raw_type_str else TransactionType.EXPENSE

                elif mapping.debit or mapping.credit:
                    debit_val = Decimal("0.00")
                    credit_val = Decimal("0.00")
                    if mapping.debit and row.get(mapping.debit):
                        raw_debit = row[mapping.debit].replace("₹", "").replace("$", "").replace(",", "").strip()
                        if raw_debit:
                            try:
                                debit_val = Decimal(raw_debit)
                            except (InvalidOperation, ValueError):
                                failed += 1
                                errors.append(ImportErrorDetail(row_number=idx, reason=f"Invalid debit amount '{raw_debit}'", data=row))
                                continue
                    if mapping.credit and row.get(mapping.credit):
                        raw_credit = row[mapping.credit].replace("₹", "").replace("$", "").replace(",", "").strip()
                        if raw_credit:
                            try:
                                credit_val = Decimal(raw_credit)
                            except (InvalidOperation, ValueError):
                                failed += 1
                                errors.append(ImportErrorDetail(row_number=idx, reason=f"Invalid credit amount '{raw_credit}'", data=row))
                                continue

                    if debit_val > Decimal("0.00"):
                        amount = debit_val
                        tx_type = TransactionType.EXPENSE
                    elif credit_val > Decimal("0.00"):
                        amount = credit_val
                        tx_type = TransactionType.INCOME
                    else:
                        failed += 1
                        errors.append(ImportErrorDetail(row_number=idx, reason="Neither debit nor credit contained a positive amount", data=row))
                        continue
                else:
                    failed += 1
                    errors.append(ImportErrorDetail(row_number=idx, reason="No mapped amount, debit, or credit column found for row", data=row))
                    continue

                if amount <= Decimal("0.00"):
                    failed += 1
                    errors.append(ImportErrorDetail(row_number=idx, reason=f"Amount must be strictly positive (got {amount})", data=row))
                    continue

                # 3. Description & Merchant
                desc = row.get(mapping.description or "", "Imported Transaction").strip() or "Imported Transaction"
                merchant = row.get(mapping.merchant or "", "").strip() or None

                # 4. Duplicate Hash Prevention
                import_hash = Transaction.generate_hash(
                    account_id=account_id,
                    tx_date=date_val,
                    amount=amount,
                    desc=desc
                )

                if import_hash in seen_hashes_in_batch:
                    if skip_duplicates:
                        skipped += 1
                        continue
                seen_hashes_in_batch.add(import_hash)

                stmt = select(Transaction).where(
                    Transaction.user_id == user_id,
                    Transaction.import_hash == import_hash
                )
                existing = (await self.session.execute(stmt)).scalar_one_or_none()
                if existing:
                    if skip_duplicates:
                        skipped += 1
                        continue

                # 5. Create Transaction
                tx = Transaction(
                    user_id=user_id,
                    account_id=account_id,
                    amount=amount,
                    currency=account.currency or "INR",
                    transaction_type=tx_type,
                    transaction_date=date_val,
                    description=desc,
                    merchant_name=merchant,
                    import_hash=import_hash,
                    is_cleared=True
                )
                self.session.add(tx)
                imported += 1

                if tx_type == TransactionType.INCOME:
                    balance_delta += amount
                elif tx_type == TransactionType.EXPENSE:
                    balance_delta -= amount

            except Exception as e:
                failed += 1
                errors.append(ImportErrorDetail(row_number=idx, reason=str(e), data=row))

        # Update account balance in batch safely
        account.current_balance = Decimal(str(account.current_balance)) + balance_delta
        await self.session.commit()

        return ImportSummaryResponse(
            total_rows_processed=len(rows),
            imported_count=imported,
            skipped_duplicates_count=skipped,
            failed_count=failed,
            errors=errors[:10],
        )
